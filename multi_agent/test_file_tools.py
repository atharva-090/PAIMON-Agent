"""Fast, deterministic tests for file_tools -- NO LLM.

Generates real fixtures (xlsx >20 rows, a small csv, a docx, a pdf) in a temp
uploads folder, then exercises every file tool and asserts the behaviour that
matters for a weak-brain agent:

  * excel_info reports the real columns + row count,
  * excel_peek shows first-10 + last-10 only when there are >20 rows,
  * excel_filter on a BAD column returns the available columns (anti-hallucination),
  * excel_search / doc_search find known content,
  * outputs respect FILE_MAX_OUTPUT (truncation).

Run:  agent/.venv/Scripts/python.exe -m multi_agent.test_file_tools
"""

from __future__ import annotations

import os
import sys
import tempfile

import file_tools as F


def _text(chunk) -> str:
    return chunk.content[0].text


def _make_fixtures(d: str) -> None:
    import pandas as pd
    from docx import Document
    from fpdf import FPDF

    # 25-row spreadsheet (>20 so peek must head/tail), with a findable row.
    rows = []
    for i in range(1, 26):
        rows.append({"id": i, "name": f"Widget-{i}", "price": i * 10,
                     "category": "premium" if i % 2 else "basic"})
    rows[12]["name"] = "Gizmo-SPECIAL"   # row id=13, a unique searchable token
    pd.DataFrame(rows).to_excel(os.path.join(d, "products.xlsx"), index=False)

    # small csv (<=20 rows -> peek shows all)
    pd.DataFrame({"city": ["Pune", "Mumbai"], "count": [3, 7]}).to_csv(
        os.path.join(d, "small.csv"), index=False)

    # docx with a known phrase
    doc = Document()
    doc.add_heading("Quarterly Report", 0)
    doc.add_paragraph("Revenue grew 12 percent in Q2 across all regions.")
    doc.save(os.path.join(d, "report.docx"))

    # text pdf with a known number
    p = FPDF()
    p.add_page()
    p.set_font("Helvetica", size=12)
    p.cell(0, 10, "Invoice total is 4200 dollars due on Friday.")
    p.output(os.path.join(d, "invoice.pdf"))


def _run() -> int:
    ok = True
    tmp = tempfile.mkdtemp(prefix="hermes_uploads_")
    # Point file_tools at the temp uploads dir.
    F.UPLOADS_DIR = tmp
    F._CACHE_DIR = os.path.join(tmp, ".cache")
    _make_fixtures(tmp)

    # --- list_uploaded_files -------------------------------------------------
    listing = _text(F.list_uploaded_files())
    for fn in ("products.xlsx", "small.csv", "report.docx", "invoice.pdf"):
        if fn not in listing:
            print(f"[FAIL] list_uploaded_files missing {fn}")
            ok = False
    print("[ok] list_uploaded_files sees all fixtures" if ok else "")

    # --- excel_info ----------------------------------------------------------
    info = _text(F.excel_info("products.xlsx"))
    if '"rows": 25' not in info or '"price"' not in info:
        print(f"[FAIL] excel_info wrong: {info[:200]}")
        ok = False
    else:
        print("[ok] excel_info: 25 rows, real columns")

    # --- excel_peek: >20 rows must head/tail ---------------------------------
    peek = _text(F.excel_peek("products.xlsx"))
    if "FIRST 10 ROWS" not in peek or "LAST 10 ROWS" not in peek \
            or "rows hidden" not in peek:
        print(f"[FAIL] excel_peek didn't head/tail a 25-row sheet:\n{peek[:300]}")
        ok = False
    elif "Widget-1" not in peek or "Widget-25" not in peek:
        print("[FAIL] excel_peek head/tail missing expected first/last rows")
        ok = False
    else:
        print("[ok] excel_peek: first-10 + last-10 on a 25-row sheet")

    # --- excel_peek: <=20 rows shows all -------------------------------------
    peek_small = _text(F.excel_peek("small.csv"))
    if "hidden" in peek_small or "Pune" not in peek_small or "Mumbai" not in peek_small:
        print(f"[FAIL] excel_peek should show ALL rows of a small csv:\n{peek_small}")
        ok = False
    else:
        print("[ok] excel_peek: small csv shown in full")

    # --- excel_filter: bad column -> available columns (anti-hallucination) --
    bad = _text(F.excel_filter("products.xlsx", "prize", ">", "100"))
    if not bad.startswith("ERROR") or "price" not in bad or "category" not in bad:
        print(f"[FAIL] bad-column filter should list real columns:\n{bad[:200]}")
        ok = False
    else:
        print("[ok] excel_filter: bad column returns the available columns")

    # --- excel_filter: numeric comparison works ------------------------------
    flt = _text(F.excel_filter("products.xlsx", "price", ">", "200"))
    # prices are 10..250; >200 -> ids 21..25 = 5 rows
    if "5 row(s)" not in flt:
        print(f"[FAIL] numeric filter wrong count:\n{flt[:200]}")
        ok = False
    else:
        print("[ok] excel_filter: numeric > comparison")

    # --- excel_search --------------------------------------------------------
    srch = _text(F.excel_search("products.xlsx", "Gizmo-SPECIAL"))
    if "Gizmo-SPECIAL" not in srch or "1 row" not in srch:
        print(f"[FAIL] excel_search didn't find the unique token:\n{srch[:200]}")
        ok = False
    else:
        print("[ok] excel_search: found the unique row")

    # --- doc_search (docx) + doc_read ----------------------------------------
    ds = _text(F.doc_search("report.docx", "Revenue"))
    if "Revenue grew 12 percent" not in ds:
        print(f"[FAIL] doc_search didn't find the docx phrase:\n{ds[:200]}")
        ok = False
    else:
        print("[ok] doc_search: found the Word phrase")
    dr = _text(F.doc_read("report.docx"))
    if "Quarterly Report" not in dr:
        print(f"[FAIL] doc_read missing heading:\n{dr[:200]}")
        ok = False
    else:
        print("[ok] doc_read: converted the docx to markdown")

    # --- doc_search (pdf) ----------------------------------------------------
    dp = _text(F.doc_search("invoice.pdf", "4200"))
    if "4200" not in dp:
        print(f"[FAIL] doc_search didn't find the pdf number:\n{dp[:200]}")
        ok = False
    else:
        print("[ok] doc_search: found the PDF number")

    # --- path traversal is rejected ------------------------------------------
    evil = _text(F.excel_info("../config.py"))
    if not evil.startswith("ERROR"):
        print(f"[FAIL] path traversal not blocked:\n{evil[:120]}")
        ok = False
    else:
        print("[ok] path traversal rejected")

    # --- truncation respects FILE_MAX_OUTPUT ---------------------------------
    saved = F._MAX
    try:
        F._MAX = 200
        big = _text(F.excel_peek("products.xlsx"))
        if "…[truncated" not in big or len(big) > 400:
            print(f"[FAIL] truncation not applied (len={len(big)})")
            ok = False
        else:
            print("[ok] output truncated to FILE_MAX_OUTPUT")
    finally:
        F._MAX = saved

    print("\n" + ("[PASS] file_tools: schema/peek/filter/search/doc all correct, "
                  "columns validated, paths sandboxed, output bounded."
                  if ok else "[FAIL] see above."))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(_run())
