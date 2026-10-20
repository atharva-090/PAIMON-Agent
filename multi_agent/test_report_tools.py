"""Fast, deterministic tests for report_tools -- NO LLM.

Exercises every renderer end-to-end against a temp reports folder and asserts the
output is a *real, valid* file of the right type (not just that a tool returned
"ok"): DOCX/PPTX/XLSX are opened back as zip/openpyxl, PNG/PDF are checked by
magic bytes, HTML by content. Also checks chart-in-document embedding, the XLSX
native chart, and that the write sandbox rejects traversal / absolute paths.

PDF goes through Microsoft Word (docx2pdf); on a machine without Word that one
check is SKIPPED (reported, not failed) so the suite still passes in CI.

Run:  agent/.venv/Scripts/python.exe -m multi_agent.test_report_tools
"""

from __future__ import annotations

import os
import tempfile
import zipfile

# Point the report sandbox at a throwaway folder BEFORE anything reads config.
_TMP = tempfile.mkdtemp(prefix="hermes_reports_")
os.environ["REPORTS_DIR"] = _TMP

import config  # noqa: E402
config.REPORTS_DIR = _TMP
import report_tools as R  # noqa: E402
R.REPORTS_DIR = _TMP


def _text(chunk) -> str:
    return chunk.content[0].text


def _saved_path(chunk) -> str:
    """Pull the saved filename out of a _saved() confirmation and return its
    absolute path in the temp reports dir (also asserts it's a SAVED message)."""
    t = _text(chunk)
    assert t.startswith("SAVED"), f"expected SAVED, got: {t[:200]}"
    # "SAVED <kind>: <name>  (..KB)"
    name = t.split(":", 1)[1].strip().split("  (")[0].strip()
    return os.path.join(_TMP, name)


def _is_zip_with(path: str, member_prefix: str) -> bool:
    if not zipfile.is_zipfile(path):
        return False
    with zipfile.ZipFile(path) as z:
        return any(n.startswith(member_prefix) for n in z.namelist())


def _run() -> int:
    ok = True

    def check(cond: bool, msg: str):
        nonlocal ok
        if cond:
            print(f"[ok] {msg}")
        else:
            print(f"[FAIL] {msg}")
            ok = False

    # --- report_rules guide --------------------------------------------------
    rules = _text(R.report_rules())
    check("write_document" in rules and "create_chart" in rules,
          "report_rules lists the generator tools")

    # --- write_document: markdown ------------------------------------------
    md = _saved_path(R.write_document("notes", "# Hi\n\n- a\n- b\n", "md"))
    check(os.path.isfile(md) and "# Hi" in open(md, encoding="utf-8").read(),
          "write_document md -> real .md file")

    # --- write_document: html (standalone, self-contained) -------------------
    html = _saved_path(R.write_document("page", "# Title\n\nHello **world**.", "html"))
    htext = open(html, encoding="utf-8").read()
    check(html.endswith(".html") and "<html" in htext.lower() and "Hello" in htext,
          "write_document html -> standalone HTML")

    # --- write_document: docx (valid OOXML zip) ------------------------------
    doc_md = "# Quarterly Report\n\nRevenue grew **12%**.\n\n| Region | Total |\n|---|---|\n| North | 120 |\n| South | 95 |\n"
    docx = _saved_path(R.write_document("q2_report", doc_md, "docx"))
    check(docx.endswith(".docx") and _is_zip_with(docx, "word/document.xml"),
          "write_document docx -> valid OOXML document")

    # --- write_document: pptx (slides) ---------------------------------------
    pptx = _saved_path(R.write_document(
        "deck", "# Slide One\n\npoint a\n\n# Slide Two\n\npoint b\n", "pptx"))
    check(pptx.endswith(".pptx") and _is_zip_with(pptx, "ppt/slides/slide"),
          "write_document pptx -> real slides")

    # --- create_chart: bar / line / pie -> valid PNGs ------------------------
    bar = _saved_path(R.create_chart(
        "sales", "bar",
        '{"labels": ["Q1","Q2","Q3"], "values": [10, 25, 17]}',
        title="Sales by Quarter", ylabel="USD"))
    check(bar.endswith(".png") and open(bar, "rb").read(4) == b"\x89PNG",
          "create_chart bar -> valid PNG")

    multi = _saved_path(R.create_chart(
        "yoy", "line",
        '{"labels": ["Q1","Q2"], "series": {"2023": [5,8], "2024": [9,12]}}',
        title="YoY"))
    check(open(multi, "rb").read(4) == b"\x89PNG", "create_chart line multi-series PNG")

    pie = _saved_path(R.create_chart(
        "share", "pie", '{"labels": ["A","B","C"], "values": [50, 30, 20]}'))
    check(open(pie, "rb").read(4) == b"\x89PNG", "create_chart pie -> valid PNG")

    # --- chart embedded in a document ----------------------------------------
    rep = _saved_path(R.write_document(
        "with_chart", "# Report\n\nSee the chart:\n\n![Sales](sales.png)\n\nDone.",
        "docx"))
    check(_is_zip_with(rep, "word/media/"),
          "write_document embeds an earlier chart PNG into the docx")

    # --- create_spreadsheet: single sheet, read back -------------------------
    xlsx = _saved_path(R.create_spreadsheet(
        "totals",
        '{"columns": ["Region","Total"], "rows": [["North",120],["South",95]]}'))
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb.active
    check(ws["A1"].value == "Region" and ws["B2"].value == 120,
          "create_spreadsheet -> xlsx with the right cells")

    # --- create_spreadsheet: multi-sheet + native chart ----------------------
    xlsx2 = _saved_path(R.create_spreadsheet(
        "book",
        '{"Sales": {"columns":["M","V"],"rows":[["Jan",10],["Feb",20]]},'
        ' "Costs": {"columns":["M","C"],"rows":[["Jan",4],["Feb",6]]}}',
        chart='{"sheet":"Sales","kind":"bar","title":"Sales"}'))
    wb2 = load_workbook(xlsx2)
    check(set(wb2.sheetnames) == {"Sales", "Costs"}, "multi-sheet workbook has both sheets")
    check(len(wb2["Sales"]._charts) >= 1, "native Excel chart added to the Sales sheet")

    # --- list_reports sees what we made --------------------------------------
    listing = _text(R.list_reports())
    check(all(n in listing for n in ("q2_report.docx", "totals.xlsx", "sales.png")),
          "list_reports lists the generated files with URLs")

    # --- write sandbox: traversal / absolute / subfolder rejected ------------
    for bad_name in ("../escape", "sub/dir/file", "/etc/passwd",
                     "C:/Windows/evil"):
        out = _text(R.write_document(bad_name, "# x", "md"))
        if not out.startswith("ERROR"):
            check(False, f"sandbox should reject {bad_name!r}: got {out[:80]}")
            break
    else:
        check(True, "write sandbox rejects traversal / absolute / subfolder names")

    # --- bad inputs are clear errors, not crashes ----------------------------
    check(_text(R.write_document("x", "body", "xls")).startswith("ERROR"),
          "unsupported fmt is a clean error")
    check(_text(R.create_chart("c", "bar", "not json")).startswith("ERROR"),
          "bad chart JSON is a clean error")

    # --- PDF via Word (skip cleanly if Word/docx2pdf unavailable) ------------
    pdf_res = _text(R.write_document("q2_report", doc_md, "pdf"))
    if pdf_res.startswith("SAVED"):
        pdf = _saved_path(R.write_document("q2_report", doc_md, "pdf"))
        check(pdf.endswith(".pdf") and open(pdf, "rb").read(5) == b"%PDF-",
              "write_document pdf -> valid PDF (via Word)")
    else:
        print(f"[skip] PDF export unavailable on this machine: {pdf_res[:120]}")

    print("\n== report_tools:", "ALL PASS ==" if ok else "FAILURES ==")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(_run())
