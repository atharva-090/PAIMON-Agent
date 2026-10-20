"""Report & document generation tools for the Hermes agent (used by the swarm's
scripter).

This is the WRITE counterpart to file_tools.py. Where file_tools READS the files
a user drops into ``data/uploads/``, report_tools WRITES real deliverables --
DOCX / PDF / XLSX / PPTX / HTML / PNG -- into ``data/reports/`` so Paimon can
hand back a finished file ("Report's ready, Traveler!").

Design (matches the rest of the project):

  * CONTENT / RENDER SPLIT. The weak ChatGPT-web brain is good at writing
    Markdown prose, bad at fiddly file-format APIs. So the agent only ever
    writes **Markdown** (or simple JSON data for sheets/charts); deterministic,
    non-AI renderers turn that into the actual binary file:
        - Markdown -> DOCX / PPTX / HTML        via pandoc (a CLI, v3.x)
        - Markdown -> PDF                       via DOCX (pandoc) then docx2pdf
                                                (Microsoft Word COM on Windows)
        - tabular JSON -> XLSX (+ native chart) via openpyxl
        - series JSON -> PNG chart              via matplotlib (Agg, no display)
      A chart PNG can be embedded in a document: create it first, then reference
      it in the Markdown as ``![caption](chart.png)`` -- pandoc resolves the name
      against the reports folder.

  * SANDBOXED WRITES. Writing is destructive, so every path is confined to a
    FLAT ``REPORTS_DIR`` -- absolute paths, ``..`` traversal and sub-folders are
    rejected, and the extension is forced to match the requested format. This is
    NOT the read-only run_python sandbox; these tools genuinely write files, but
    only ever inside the one output folder.

Tools: report_rules (read first), write_document, create_chart,
create_spreadsheet, list_reports.
"""

from __future__ import annotations

import io
import json
import os
import re
import subprocess
import tempfile
from typing import Optional

from agentscope.message import TextBlock
from agentscope.tool import FunctionTool, ToolChunk

import config

REPORTS_DIR = config.REPORTS_DIR
PANDOC = config.PANDOC_BIN
_MAX = config.FILE_MAX_OUTPUT

# Which extensions each format produces. The format the agent names is the
# source of truth; we force the filename's extension to match so a "report.txt"
# asked to be a docx still lands as "report.docx".
_DOC_FORMATS = {"docx": ".docx", "pdf": ".pdf", "html": ".html",
                "pptx": ".pptx", "md": ".md"}
_CHART_KINDS = {"bar", "line", "pie", "barh"}


# ===========================================================================
#  helpers
# ===========================================================================
def _chunk(text: str) -> ToolChunk:
    if len(text) > _MAX:
        text = text[:_MAX] + f"\n…[truncated — {len(text)} chars]"
    return ToolChunk(content=[TextBlock(type="text", text=text)])


def _ok(obj) -> ToolChunk:
    return _chunk(obj if isinstance(obj, str)
                  else json.dumps(obj, ensure_ascii=False, indent=2, default=str))


def _err(msg: str) -> ToolChunk:
    return _chunk(f"ERROR: {msg}")


def _resolve_out(filename: str, force_ext: str) -> str:
    """Resolve ``filename`` to an absolute path inside REPORTS_DIR, forcing the
    extension to ``force_ext``. Rejects absolute paths, traversal, and
    sub-folders so the output area stays a flat, sandboxed directory.

    The agent passes a bare name like 'q2_report' or 'q2_report.docx'; we strip a
    leading 'reports/' for convenience, refuse anything with a path separator,
    and make the extension match the format being written."""
    name = (filename or "").strip().replace("\\", "/").strip()
    name = name.lstrip("/")
    for prefix in ("reports/", "data/reports/"):
        if name.lower().startswith(prefix):
            name = name[len(prefix):]
    if not name:
        raise ValueError("filename is empty")
    if "/" in name or name.startswith("."):
        raise ValueError(
            f"use a plain filename like 'report{force_ext}' -- no folders, no "
            "leading dot, no path")
    stem = os.path.splitext(name)[0]
    if not stem:
        raise ValueError("filename has no name part")
    name = stem + force_ext
    base = os.path.realpath(REPORTS_DIR)
    full = os.path.realpath(os.path.join(base, name))
    if full != base and not full.startswith(base + os.sep):
        raise ValueError(f"'{filename}' resolves outside the reports folder")
    os.makedirs(base, exist_ok=True)
    return full


def _download_url(full: str) -> str:
    return f"/api/reports/{os.path.basename(full)}"


def _saved(full: str, kind: str, extra: str = "") -> ToolChunk:
    """Standard 'file written' confirmation: what was made, how big, and the URL
    the user can download it from. The agent should relay the download link."""
    size_kb = round(os.path.getsize(full) / 1024, 1)
    msg = (f"SAVED {kind}: {os.path.basename(full)}  ({size_kb} KB)\n"
           f"Download: {_download_url(full)}")
    if extra:
        msg += f"\n{extra}"
    return _ok(msg)


def _parse_json_arg(value, what: str):
    """Accept either a JSON string (what the brain usually emits) or an
    already-decoded dict/list. Raises ValueError with a clear message."""
    if isinstance(value, (dict, list)):
        return value
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{what} is empty -- pass it as a JSON object")
    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{what} is not valid JSON: {exc}")


# ===========================================================================
#  rendering primitives (deterministic, no AI)
# ===========================================================================
def _run_pandoc(md_text: str, out_path: str, extra_args=()) -> None:
    """Render Markdown text to ``out_path`` via the pandoc CLI. Images referenced
    as ``![](name.png)`` resolve against REPORTS_DIR (so a chart made earlier is
    embedded). Raises RuntimeError on a pandoc failure."""
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False,
                                     encoding="utf-8") as tf:
        tf.write(md_text)
        src = tf.name
    try:
        cmd = [PANDOC, src, "-f", "markdown", "-o", out_path,
               "--resource-path", REPORTS_DIR, *extra_args]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
        if proc.returncode != 0:
            raise RuntimeError(
                f"pandoc failed (exit {proc.returncode}): "
                f"{(proc.stderr or proc.stdout).strip()[:500]}")
    except FileNotFoundError:
        raise RuntimeError(
            f"pandoc not found (looked for '{PANDOC}'). Install pandoc or set "
            "PANDOC_BIN; DOCX/PPTX/HTML/PDF need it.")
    finally:
        try:
            os.unlink(src)
        except OSError:
            pass


def _docx_to_pdf(docx_path: str, pdf_path: str) -> None:
    """Convert a .docx to .pdf using Microsoft Word via docx2pdf (COM). Word's
    COM wants the calling thread initialised, and docx2pdf prints a tqdm bar we
    don't want in tool output, so we init COM defensively and swallow its
    stdout/stderr. Raises RuntimeError with an actionable message if Word/the
    converter isn't usable."""
    did_co = False
    try:
        import pythoncom  # ships with pywin32 (a docx2pdf dependency on Windows)
        pythoncom.CoInitialize()
        did_co = True
    except Exception:
        pass  # convert() still works on the main thread without this
    try:
        from docx2pdf import convert
    except ImportError:
        raise RuntimeError(
            "docx2pdf isn't installed, so PDF export is unavailable. Ask for "
            "DOCX or HTML instead, or `pip install docx2pdf` (needs Microsoft "
            "Word).")
    try:
        sink = io.StringIO()
        from contextlib import redirect_stdout, redirect_stderr
        with redirect_stdout(sink), redirect_stderr(sink):
            convert(docx_path, pdf_path)
    except Exception as exc:
        raise RuntimeError(
            f"Word PDF conversion failed ({type(exc).__name__}: {exc}). This path "
            "needs Microsoft Word installed; try DOCX or HTML instead.")
    finally:
        if did_co:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass
    if not os.path.isfile(pdf_path):
        raise RuntimeError(
            "PDF conversion produced no file (is Microsoft Word installed and "
            "licensed?). Try DOCX or HTML instead.")


def _render_chart(kind: str, data: dict, out_path: str, title: str,
                  xlabel: str, ylabel: str) -> None:
    """Render a chart PNG with matplotlib (Agg backend -- no display needed).

    ``data`` shapes:
      bar/line/barh: {"labels": [...], "values": [...]}  (single series), or
                     {"labels": [...], "series": {"A": [...], "B": [...]}}
      pie:           {"labels": [...], "values": [...]}
    """
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    labels = data.get("labels")
    if not labels or not isinstance(labels, list):
        raise ValueError("chart data needs a non-empty 'labels' list")

    fig, ax = plt.subplots(figsize=(7, 4.2))
    try:
        if kind == "pie":
            values = data.get("values")
            if not values:
                raise ValueError("a pie chart needs a 'values' list")
            ax.pie([float(v) for v in values], labels=[str(x) for x in labels],
                   autopct="%1.1f%%", startangle=90)
            ax.axis("equal")
        else:
            series = data.get("series")
            if not series:
                values = data.get("values")
                if not values:
                    raise ValueError(
                        "chart data needs 'values' (single series) or 'series' "
                        "(a {name: [...]} object for multiple series)")
                series = {ylabel or "value": values}
            x = range(len(labels))
            n = len(series)
            for i, (name, ys) in enumerate(series.items()):
                ys = [float(v) for v in ys]
                if kind == "bar":
                    width = 0.8 / max(n, 1)
                    ax.bar([p + i * width for p in x], ys, width=width, label=str(name))
                elif kind == "barh":
                    ax.barh([str(l) for l in labels], ys, label=str(name))
                else:  # line
                    ax.plot([str(l) for l in labels], ys, marker="o", label=str(name))
            if kind == "bar":
                ax.set_xticks([p + 0.4 - (0.8 / max(n, 1)) / 2 for p in x])
                ax.set_xticklabels([str(l) for l in labels], rotation=0)
            if n > 1:
                ax.legend()
            if xlabel:
                ax.set_xlabel(xlabel)
            if ylabel:
                ax.set_ylabel(ylabel)
        if title:
            ax.set_title(title)
        fig.tight_layout()
        fig.savefig(out_path, dpi=130)
    finally:
        plt.close(fig)


def _write_xlsx(spec: dict, out_path: str, chart: Optional[dict]) -> list:
    """Write an .xlsx from a sheet spec via openpyxl. Returns the sheet names.

    ``spec`` is either a single sheet {"columns": [...], "rows": [[...], ...]}
    or multiple {"SheetName": {"columns": [...], "rows": [...]}, ...}.
    ``chart`` (optional): {"sheet": name, "kind": "bar"|"line"|"pie",
    "title": str} -- a native Excel chart of the numeric columns vs the first
    column as categories."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    # Normalise to {sheet_name: {columns, rows}}.
    if "columns" in spec or "rows" in spec:
        sheets = {"Sheet1": spec}
    else:
        sheets = spec
    if not sheets:
        raise ValueError("no sheet data provided")

    wb = Workbook()
    wb.remove(wb.active)
    written = []
    for sname, sdata in sheets.items():
        cols = sdata.get("columns") or []
        rows = sdata.get("rows") or []
        ws = wb.create_sheet(title=str(sname)[:31])  # Excel sheet-name limit
        if cols:
            ws.append([str(c) for c in cols])
        for r in rows:
            ws.append(list(r))
        written.append(ws.title)

    if chart:
        cname = chart.get("sheet") or written[0]
        if cname not in wb.sheetnames:
            raise ValueError(f"chart sheet '{cname}' not found; sheets: {written}")
        ws = wb[cname]
        nrows, ncols = ws.max_row, ws.max_column
        if nrows >= 2 and ncols >= 2:
            kind = (chart.get("kind") or "bar").lower()
            ch = {"line": LineChart, "pie": PieChart}.get(kind, BarChart)()
            ch.title = chart.get("title") or None
            data = Reference(ws, min_col=2, min_row=1, max_col=ncols, max_row=nrows)
            cats = Reference(ws, min_col=1, min_row=2, max_row=nrows)
            ch.add_data(data, titles_from_data=True)
            ch.set_categories(cats)
            ws.add_chart(ch, f"{chr(ord('A') + ncols + 1)}2")

    wb.save(out_path)
    return written


# ===========================================================================
#  guide tool (read first -- like python_coding_rules for code mode)
# ===========================================================================
def report_rules() -> ToolChunk:
    """Read this FIRST before generating any document/report/spreadsheet/chart.
    It tells you exactly which tool to call and the data shapes they expect, so
    you produce a real file instead of pasting content into chat.
    """
    return _ok(
        "REPORT GENERATION -- how to produce a real downloadable file\n"
        "===========================================================\n"
        "You WRITE Markdown (or simple JSON data); deterministic renderers turn "
        "it into a real DOCX/PDF/XLSX/PPTX/HTML/PNG saved in the reports folder "
        "and downloadable at /api/reports/<filename>. NEVER paste a whole report "
        "into the chat -- always call a tool so the user gets a file.\n\n"
        "TOOLS:\n"
        "1. write_document(filename, markdown, fmt) -- the main one. fmt is one "
        "of: docx, pdf, html, pptx, md.\n"
        "   * Write normal Markdown: '# Title', '## Heading', **bold**, '- bullet',"
        " '| a | b |' tables, etc.\n"
        "   * PPTX: each top-level '# Heading' starts a new slide.\n"
        "   * To put a CHART in the document: first call create_chart(...) to make"
        " 'sales.png', THEN reference it in the markdown as '![Sales](sales.png)'."
        " Use just the filename (it lives in the same folder).\n"
        "2. create_chart(filename, kind, data, title, xlabel, ylabel) -- makes a "
        "PNG. kind: bar, line, pie, barh. data is a JSON object:\n"
        "     {\"labels\": [\"Q1\",\"Q2\",\"Q3\"], \"values\": [10, 25, 17]}\n"
        "   or for several series:\n"
        "     {\"labels\": [\"Q1\",\"Q2\"], \"series\": {\"2023\": [5,8], \"2024\": [9,12]}}\n"
        "3. create_spreadsheet(filename, data, chart) -- makes an XLSX. data is a "
        "JSON object, either one sheet:\n"
        "     {\"columns\": [\"Region\",\"Total\"], \"rows\": [[\"North\",120],[\"South\",95]]}\n"
        "   or several: {\"Sales\": {\"columns\":[...],\"rows\":[...]}, \"Costs\": {...}}.\n"
        "   chart (optional JSON) adds a native Excel chart: "
        "{\"sheet\":\"Sheet1\",\"kind\":\"bar\",\"title\":\"Totals\"} (first column = "
        "categories, the numeric columns = series).\n"
        "4. list_reports() -- see what you've already generated.\n\n"
        "WORKFLOW: if a report draws on an uploaded file, read it first with the "
        "file tools (excel_*/query_data/doc_*), compute the real numbers, THEN "
        "write the document/spreadsheet from those numbers -- never invent data. "
        "Filenames are plain (no folders): 'q2_summary.docx'. After saving, tell "
        "the user the download link the tool returned."
    )


# ===========================================================================
#  generation tools (write-capable, sandboxed to REPORTS_DIR)
# ===========================================================================
def write_document(filename: str, markdown: str, fmt: str = "docx") -> ToolChunk:
    """Generate a real document FILE from Markdown and save it to the reports
    folder for download. This is how you deliver a written report/letter/memo/
    summary as an actual DOCX, PDF, HTML, PowerPoint, or Markdown file instead of
    pasting text into the chat.

    Write ordinary Markdown for the body: '# Title', '## Section', **bold**,
    bullet '- ' lists, numbered lists, and '| col | col |' tables all work. For a
    PowerPoint (fmt='pptx'), each top-level '# Heading' begins a new slide. To
    embed a chart you made with create_chart, reference it in the markdown as
    '![caption](chart.png)' using just the filename.

    Args:
        filename (str): Output name, e.g. 'q2_summary' or 'q2_summary.docx'
            (the extension is set to match fmt automatically; no folders).
        markdown (str): The document body as Markdown.
        fmt (str): One of docx, pdf, html, pptx, md. PDF is produced via Word
            (Windows); if Word is unavailable, use docx or html.
    """
    try:
        fmt = (fmt or "docx").strip().lower().lstrip(".")
        if fmt not in _DOC_FORMATS:
            return _err(f"unsupported fmt '{fmt}'. Use one of: "
                        f"{', '.join(_DOC_FORMATS)}.")
        if not (markdown or "").strip():
            return _err("markdown is empty -- pass the document body as Markdown.")
        out = _resolve_out(filename, _DOC_FORMATS[fmt])

        if fmt == "md":
            with open(out, "w", encoding="utf-8") as fh:
                fh.write(markdown)
        elif fmt == "html":
            _run_pandoc(markdown, out, extra_args=("--standalone",
                                                   "--embed-resources"))
        elif fmt in ("docx", "pptx"):
            _run_pandoc(markdown, out)
        elif fmt == "pdf":
            # Markdown -> DOCX (pandoc) -> PDF (Word). The intermediate DOCX is a
            # throwaway temp file; only the .pdf lands in the reports folder.
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
                tmp_docx = tf.name
            try:
                _run_pandoc(markdown, tmp_docx)
                _docx_to_pdf(tmp_docx, out)
            finally:
                try:
                    os.unlink(tmp_docx)
                except OSError:
                    pass
        return _saved(out, f"{fmt.upper()} document")
    except (ValueError, RuntimeError) as exc:
        return _err(str(exc))
    except Exception as exc:
        return _err(f"could not write document: {type(exc).__name__}: {exc}")


def create_chart(filename: str, kind: str, data: str, title: str = "",
                 xlabel: str = "", ylabel: str = "") -> ToolChunk:
    """Generate a chart as a PNG image saved to the reports folder. Use it to
    visualize numbers -- on its own, or (more often) to embed in a document:
    create the chart first, then reference it in write_document's markdown as
    '![caption](filename.png)'.

    Args:
        filename (str): Output name, e.g. 'sales' or 'sales.png' (no folders).
        kind (str): bar, line, pie, or barh (horizontal bar).
        data (str): A JSON object describing the data. Single series:
            {"labels": ["Q1","Q2","Q3"], "values": [10, 25, 17]}.
            Multiple series (bar/line): {"labels": ["Q1","Q2"],
            "series": {"2023": [5,8], "2024": [9,12]}}. A pie chart uses
            labels + values.
        title (str): Optional chart title.
        xlabel (str): Optional x-axis label.
        ylabel (str): Optional y-axis label (also names a single series).
    """
    try:
        kind = (kind or "bar").strip().lower()
        if kind not in _CHART_KINDS:
            return _err(f"unsupported kind '{kind}'. Use one of: "
                        f"{', '.join(sorted(_CHART_KINDS))}.")
        spec = _parse_json_arg(data, "chart data")
        if not isinstance(spec, dict):
            return _err("chart data must be a JSON object with 'labels' and "
                        "'values' (or 'series').")
        out = _resolve_out(filename, ".png")
        _render_chart(kind, spec, out, title, xlabel, ylabel)
        return _saved(out, f"{kind} chart",
                      extra=f"Embed it in a document with: ![{title or 'chart'}]"
                            f"({os.path.basename(out)})")
    except (ValueError, RuntimeError) as exc:
        return _err(str(exc))
    except Exception as exc:
        return _err(f"could not create chart: {type(exc).__name__}: {exc}")


def create_spreadsheet(filename: str, data: str, chart: str = "") -> ToolChunk:
    """Generate an Excel (.xlsx) spreadsheet from tabular data and save it to the
    reports folder. Use it to deliver computed results as a real workbook --
    optionally with a native Excel chart.

    Args:
        filename (str): Output name, e.g. 'q2_totals' or 'q2_totals.xlsx'.
        data (str): A JSON object. One sheet:
            {"columns": ["Region","Total"], "rows": [["North",120],["South",95]]}.
            Several sheets: {"Sales": {"columns":[...],"rows":[...]},
            "Costs": {"columns":[...],"rows":[...]}}.
        chart (str): Optional JSON for a native chart, e.g.
            {"sheet": "Sheet1", "kind": "bar", "title": "Totals"} -- the first
            column is the categories and the numeric columns are the series.
    """
    try:
        spec = _parse_json_arg(data, "spreadsheet data")
        if not isinstance(spec, dict):
            return _err("spreadsheet data must be a JSON object (see report_rules).")
        chart_spec = None
        if (chart or "").strip():
            chart_spec = _parse_json_arg(chart, "chart spec")
        out = _resolve_out(filename, ".xlsx")
        sheets = _write_xlsx(spec, out, chart_spec)
        return _saved(out, "XLSX spreadsheet",
                      extra=f"sheets: {', '.join(sheets)}")
    except (ValueError, RuntimeError) as exc:
        return _err(str(exc))
    except Exception as exc:
        return _err(f"could not write spreadsheet: {type(exc).__name__}: {exc}")


def _scan_reports() -> list:
    """Raw listing of generated reports (newest first) -- the data the API and
    the list_reports tool both render."""
    base = config.REPORTS_DIR
    if not os.path.isdir(base):
        return []
    out = []
    for name in os.listdir(base):
        full = os.path.join(base, name)
        # Skip dotfiles and Office owner/lock files ("~$name.docx" that Word
        # leaves behind when a doc is open) so they don't show up as fake reports.
        if name.startswith((".", "~$")) or not os.path.isfile(full):
            continue
        out.append({
            "filename": name,
            "type": os.path.splitext(name)[1].lstrip(".").lower() or "?",
            "size_kb": round(os.path.getsize(full) / 1024, 1),
            "modified": os.path.getmtime(full),
            "url": f"/api/reports/{name}",
        })
    out.sort(key=lambda r: r["modified"], reverse=True)
    return out


def list_reports() -> ToolChunk:
    """List the report files already generated in the reports folder (newest
    first), with their download URLs. Use it to check what you've produced or to
    re-share a download link.
    """
    try:
        rows = _scan_reports()
        if not rows:
            return _ok(f"No reports generated yet (folder: {REPORTS_DIR}).")
        view = [{k: r[k] for k in ("filename", "type", "size_kb", "url")}
                for r in rows]
        return _ok({"reports_dir": REPORTS_DIR, "reports": view})
    except Exception as exc:
        return _err(f"could not list reports: {type(exc).__name__}: {exc}")


# ===========================================================================
#  factory
# ===========================================================================
# report_rules + list_reports are read-only; the three generators write files
# (sandboxed to REPORTS_DIR), so they are NOT marked read-only.
_READONLY = {report_rules, list_reports}
_REPORT_FUNCS = [
    report_rules,
    write_document, create_chart, create_spreadsheet,
    list_reports,
]


def build_tools() -> list:
    """Return the report tools as AgentScope FunctionTools."""
    return [FunctionTool(f, is_read_only=(f in _READONLY)) for f in _REPORT_FUNCS]
